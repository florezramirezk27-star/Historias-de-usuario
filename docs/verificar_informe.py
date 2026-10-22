# -*- coding: utf-8 -*-
"""Revisa artefactos de Markdown que могли quedar en el .docx."""
from docx import Document

d = Document(r"docs\Sistema_Trabajos_de_Grado_UTS.docx")
marcas = ["```", "|---", "**", "## ", "~~", "| ", "]("]
sospechosos = []
for p in d.paragraphs:
    for m in marcas:
        if m in p.text:
            sospechosos.append(("P", m, p.text[:110]))
            break
for i, tb in enumerate(d.tables):
    for r in tb.rows:
        for c in r.cells:
            for m in marcas:
                if m in c.text:
                    sospechosos.append(("T%d" % i, m, c.text[:110]))
                    break
print("Artefactos:", len(sospechosos))
for x in sospechosos[:30]:
    print(x)
